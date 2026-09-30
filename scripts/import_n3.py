"""Import the supplied N3 workbook and add Sino-Vietnamese readings.

Run: python scripts/import_n3.py [workbook.xlsx]
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import urllib.request
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = Path(r"D:\học liệu\2026.1\jlpt\[VTI Mirai share] 140 - 880 TỪ VỰNG MIMIKARA N3 MINATO.xlsx")
UNIHAN_URL = "https://www.unicode.org/Public/17.0.0/ucd/Unihan.zip"
KANJI_URL = "https://raw.githubusercontent.com/NVL4826/kanji-data-hanviet/master/kanji-jouyou.json"
OVERRIDES_PATH = ROOT / "scripts" / "n3_hanviet_overrides.json"
CHAR_OVERRIDES = {
    "渇": "khát", "讀": "độc", "伝": "truyền", "広": "quảng",
    "死": "tử", "数": "số", "番": "phiên", "営": "doanh",
    "実": "thực", "激": "kích", "条": "điều", "案": "án",
    "参": "tham", "炊": "xuy", "徴": "trưng", "籍": "tịch",
    "東": "đông", "儘": "tận", "詫": "sá", "読": "độc",
    "行": "hành", "調": "điều", "重": "trọng", "長": "trưởng",
    "句": "câu", "期": "kỳ", "比": "tỷ", "理": "lý",
    "予": "dự", "時": "thời", "違": "vi",
}
READING_OVERRIDES = {90: "さくしゃ", 178: "すれちがう", 750: "すませる・すます"}


def is_han(char: str) -> bool:
    return "CJK UNIFIED IDEOGRAPH" in unicodedata.name(char, "")


def unihan_readings() -> dict[str, list[str]]:
    with urllib.request.urlopen(UNIHAN_URL, timeout=60) as response:
        archive = zipfile.ZipFile(BytesIO(response.read()))
    result: dict[str, list[str]] = {}
    for line in archive.read("Unihan_Readings.txt").decode("utf-8").splitlines():
        if "\tkVietnamese\t" not in line:
            continue
        code, _, value = line.split("\t")
        result[chr(int(code[2:], 16))] = value.split()
    return result


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    columns = [i for i, label in enumerate(rows[2]) if label == "STT"]
    readings = unihan_readings()
    with urllib.request.urlopen(KANJI_URL, timeout=60) as response:
        kanji_data = json.load(response)
    overrides = json.loads(OVERRIDES_PATH.read_text(encoding="utf-8")) if OVERRIDES_PATH.exists() else {}
    result = {}
    ambiguous = []
    missing = []
    for row in rows[3:]:
        for i in columns:
            number = row[i]
            if not isinstance(number, int):
                continue
            word = str(row[i + 1]).strip()
            has_reading = rows[2][i + 2] == "Hiragana"
            reading = READING_OVERRIDES.get(number, row[i + 2] if has_reading else word)
            meaning = row[i + 3] if has_reading else row[i + 2]
            han_chars = [char for char in word if is_han(char)]
            han_viet = None
            if han_chars:
                parts = []
                for char in han_chars:
                    # Unihan's first kVietnamese value is often a rare variant
                    # (e.g. 女=nữa); prefer the curated common reading, while
                    # retaining Unihan as the fallback and audit source.
                    options = [CHAR_OVERRIDES[char]] if char in CHAR_OVERRIDES else (kanji_data.get(char, {}).get("han_viet", []) or readings.get(char, []) or [])
                    if not options:
                        missing.append((number, char))
                    if len(options) > 1:
                        ambiguous.append((number, char, options))
                    parts.append(options[0] if options else "?")
                han_viet = " ".join(parts).upper()
            if str(number) in overrides:
                han_viet = overrides[str(number)]
            result[number] = {
                "id": f"n3:{number}",
                "level": "N3",
                "order": number,
                "headword": word,
                "reading": str(reading).strip(),
                "hanViet": han_viet,
                "meanings": [part.strip() for part in re.split(r"[,;；]", str(meaning)) if part.strip()],
            }
    if sorted(result) != list(range(1, 881)):
        raise ValueError("N3 numbers must be exactly 1..880")
    output = ROOT / "src" / "data" / "n3.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps([result[i] for i in sorted(result)], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(result)} N3 entries to {output}")
    print(f"Missing Han Viet characters: {missing!a}")
    print(f"Ambiguous character readings: {len(ambiguous)}")
    (ROOT / "scripts" / "n3_hanviet_review.json").write_text(json.dumps(ambiguous, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
